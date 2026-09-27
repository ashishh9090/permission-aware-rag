import io
import uuid
import hashlib
import time
from typing import List, Tuple, Dict, Any
from sqlalchemy.orm import Session
from fastapi import UploadFile, HTTPException, status
import pypdf
import docx

from backend.app.models.document import Document
from backend.app.models.chunk import Chunk
from backend.app.models.user import User
from backend.app.services.embedding_service import embedding_service
from backend.app.services.audit_service import audit_service

class IngestionService:
    MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB

    @staticmethod
    def extract_text(filename: str, content_bytes: bytes) -> str:
        ext = filename.split(".")[-1].lower()
        if ext in ["txt", "md"]:
            return content_bytes.decode("utf-8", errors="ignore")
        elif ext == "pdf":
            try:
                reader = pypdf.PdfReader(io.BytesIO(content_bytes))
                text_parts = []
                for i, page in enumerate(reader.pages):
                    page_text = page.extract_text() or ""
                    if page_text.strip():
                        text_parts.append(f"[Page {i+1}]\n{page_text}")
                if text_parts:
                    return "\n\n".join(text_parts)
            except Exception:
                pass
            return content_bytes.decode("utf-8", errors="ignore")
        elif ext == "docx":
            try:
                doc = docx.Document(io.BytesIO(content_bytes))
                text_parts = [p.text for p in doc.paragraphs if p.text.strip()]
                if text_parts:
                    return "\n\n".join(text_parts)
            except Exception:
                pass
            return content_bytes.decode("utf-8", errors="ignore")
        else:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Unsupported document format: .{ext}. Allowed: .txt, .md, .pdf, .docx"
            )

    @staticmethod
    def chunk_text(text: str, chunk_size: int = 500, overlap: int = 100) -> List[Tuple[str, str]]:
        """
        Splits text into chunks, preserving section references if present.
        Returns a list of tuples: (chunk_text, source_reference)
        """
        paragraphs = text.split("\n\n")
        chunks = []
        current_chunk = []
        current_len = 0
        current_ref = "Section 1"

        for p in paragraphs:
            p_strip = p.strip()
            if not p_strip:
                continue

            if p_strip.startswith("[Page ") and "]" in p_strip:
                current_ref = p_strip[:p_strip.find("]") + 1]

            words = p_strip.split()
            for word in words:
                current_chunk.append(word)
                current_len += len(word) + 1

                if current_len >= chunk_size:
                    chunk_str = " ".join(current_chunk)
                    chunks.append((chunk_str, current_ref))
                    # Overlap retention
                    overlap_words = current_chunk[-(overlap // 6):] if len(current_chunk) > (overlap // 6) else []
                    current_chunk = list(overlap_words)
                    current_len = sum(len(w) + 1 for w in current_chunk)

        if current_chunk:
            chunk_str = " ".join(current_chunk)
            chunks.append((chunk_str, current_ref))

        return chunks if chunks else [(text[:chunk_size], "Full Text")]

    @classmethod
    def ingest_document(
        cls,
        db: Session,
        user: User,
        document_name: str,
        document_type: str,
        content_bytes: bytes,
        classification: str,
        allowed_roles: List[str],
        allowed_users: List[str],
        allowed_groups: List[str],
        summary: str = "",
        request_id: str = None
    ) -> Document:
        start_time = time.time()
        req_id = request_id or f"req-{uuid.uuid4().hex[:8]}"

        # 1. Validation
        if len(content_bytes) > cls.MAX_FILE_SIZE:
            raise HTTPException(
                status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                detail=f"File exceeds maximum allowed size of {cls.MAX_FILE_SIZE // (1024*1024)}MB"
            )

        if not allowed_roles and not allowed_users and not allowed_groups:
            # Default-deny: Rejects upload without valid access policy
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Security policy requires at least one allowed role, group, or user."
            )

        # 2. Text Extraction & Hashing
        text = cls.extract_text(document_name, content_bytes)
        content_hash = hashlib.sha256(content_bytes).hexdigest()

        # 3. Create Document Record
        doc_id = f"doc_{uuid.uuid4().hex[:12]}"
        document = Document(
            document_id=doc_id,
            tenant_id=user.tenant_id,  # Strictly server-side user tenant
            owner_id=user.id,
            document_name=document_name,
            document_type=document_type.lower(),
            classification=classification,
            allowed_roles=allowed_roles,
            allowed_users=allowed_users,
            allowed_groups=allowed_groups,
            access_policy_version=1,
            ingestion_status="indexed",
            content_hash=content_hash,
            summary=summary or text[:200],
            metadata_json={"raw_length": len(text)}
        )
        db.add(document)
        db.flush()

        # 4. Text Chunking
        raw_chunks = cls.chunk_text(text)
        created_chunk_ids = []

        for idx, (chunk_str, source_ref) in enumerate(raw_chunks):
            chunk_id = f"chk_{uuid.uuid4().hex[:12]}"
            emb = embedding_service.embed_text(chunk_str)

            chunk = Chunk(
                chunk_id=chunk_id,
                document_id=doc_id,
                tenant_id=user.tenant_id,  # Mandatory tenant isolation
                chunk_text=chunk_str,
                embedding=emb,
                classification=classification,
                allowed_roles=allowed_roles,
                allowed_users=allowed_users,
                allowed_groups=allowed_groups,
                access_policy_version=1,
                source_reference=f"{source_ref}, Chunk {idx+1}",
                chunk_index=idx,
                metadata_json={"token_count": len(chunk_str.split())}
            )
            db.add(chunk)
            created_chunk_ids.append(chunk_id)

        db.commit()
        db.refresh(document)

        latency_ms = (time.time() - start_time) * 1000.0

        # 5. Audit Logging
        audit_service.log_event(
            db=db,
            request_id=req_id,
            tenant_id=user.tenant_id,
            user_id=user.id,
            action="document_ingestion",
            policy_version=1,
            query_text=f"Ingest {document_name} ({len(raw_chunks)} chunks)",
            retrieved_chunk_ids=[],
            authorized_chunk_ids=created_chunk_ids,
            denied_count=0,
            doc_ids_sent=[doc_id],
            decision="PERMIT",
            outcome="SUCCESS",
            latency_ms=latency_ms,
            details={
                "document_name": document_name,
                "classification": classification,
                "chunks_count": len(raw_chunks),
                "content_hash": content_hash
            }
        )

        return document

    @staticmethod
    def update_document_permissions(
        db: Session,
        document_id: str,
        user: User,
        classification: str = None,
        allowed_roles: List[str] = None,
        allowed_users: List[str] = None,
        allowed_groups: List[str] = None,
        request_id: str = None
    ) -> Document:
        """
        Updates document permissions and propagates atomically to existing chunks.
        Increments access_policy_version.
        """
        req_id = request_id or f"req-{uuid.uuid4().hex[:8]}"
        doc = db.query(Document).filter(
            Document.document_id == document_id,
            Document.tenant_id == user.tenant_id
        ).first()

        if not doc:
            raise HTTPException(status_code=404, detail="Document not found")

        # Only Owner or Tenant Admin can modify permissions
        user_roles = set(user.roles or [])
        if user.id != doc.owner_id and "Tenant Admin" not in user_roles and "Super Admin" not in user_roles:
            raise HTTPException(
                status_code=403,
                detail="Forbidden: Insufficient privileges to update document permissions"
            )

        # Increment access policy version
        doc.access_policy_version += 1
        new_version = doc.access_policy_version

        if classification is not None:
            doc.classification = classification
        if allowed_roles is not None:
            doc.allowed_roles = allowed_roles
        if allowed_users is not None:
            doc.allowed_users = allowed_users
        if allowed_groups is not None:
            doc.allowed_groups = allowed_groups

        # Atomic cascade propagation to all chunks
        chunks = db.query(Chunk).filter(
            Chunk.document_id == document_id,
            Chunk.tenant_id == user.tenant_id
        ).all()

        for chk in chunks:
            chk.access_policy_version = new_version
            if classification is not None:
                chk.classification = classification
            if allowed_roles is not None:
                chk.allowed_roles = allowed_roles
            if allowed_users is not None:
                chk.allowed_users = allowed_users
            if allowed_groups is not None:
                chk.allowed_groups = allowed_groups

        db.commit()
        db.refresh(doc)

        audit_service.log_event(
            db=db,
            request_id=req_id,
            tenant_id=user.tenant_id,
            user_id=user.id,
            action="permission_update",
            policy_version=new_version,
            query_text=f"Update permissions for {doc.document_name}",
            retrieved_chunk_ids=[],
            authorized_chunk_ids=[c.chunk_id for c in chunks],
            denied_count=0,
            doc_ids_sent=[doc.document_id],
            decision="PERMIT",
            outcome="SUCCESS",
            latency_ms=10.0,
            details={
                "document_id": document_id,
                "new_version": new_version,
                "new_classification": doc.classification,
                "updated_chunks": len(chunks)
            }
        )

        return doc

ingestion_service = IngestionService()
